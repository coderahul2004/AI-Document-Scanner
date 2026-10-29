import sys
try:
    __import__("pysqlite3")
    sys.modules["sqlite3"] = sys.modules.pop("pysqlite3")
except ImportError:
    pass
import os
import time
import asyncio
import pandas as pd
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_community.document_loaders import PyPDFLoader
from docx import Document as DocxDocument
from config import GOOGLE_API_KEY, EMBEDDING_MODEL, CHROMA_DIR, COLLECTION_NAME
embeddings = None  
def get_embeddings():
    global embeddings
    if embeddings is None:
        try:
            asyncio.get_event_loop()
        except RuntimeError:
            asyncio.set_event_loop(asyncio.new_event_loop())
        embeddings = GoogleGenerativeAIEmbeddings(
            model=EMBEDDING_MODEL,
            google_api_key=GOOGLE_API_KEY
        )
    return embeddings
def get_vectorstore(collection_name=None):
    return Chroma(
        collection_name=collection_name or COLLECTION_NAME,
        embedding_function=get_embeddings(),
        persist_directory=CHROMA_DIR
    )
def table_to_documents(data, file_name, file_type, sheet_name=""):
    rows = []
    for number, (_, row) in enumerate(data.iterrows(), start=2):  
        text = "\n".join(
            f"{column}: {row[column]}"
            for column in data.columns
            if pd.notna(row[column])
        )
        if text.strip():
            rows.append((number, f"Row {number}:\n{text}"))
    documents = []
    for i in range(0, len(rows), 10):
        group = rows[i:i + 10]
        metadata = {"source": file_name, "file_type": file_type, "row": group[0][0]}
        if sheet_name:
            metadata["sheet"] = str(sheet_name)
        documents.append(Document(
            page_content="\n\n".join(item[1] for item in group),
            metadata=metadata
        ))
    return documents
def load_file(path):
    extension = os.path.splitext(path)[1].lower()
    file_name = os.path.basename(path)
    documents = []
    if extension == ".pdf":
        pages = PyPDFLoader(path).load()
        for page in pages:
            if page.page_content.strip():
                page.metadata.update({
                    "source": file_name,
                    "file_type": "pdf",
                    "page": page.metadata.get("page", 0) + 1  
                })
                documents.append(page)
    elif extension == ".docx":
        word_file = DocxDocument(path)
        lines = [p.text for p in word_file.paragraphs if p.text.strip()]
        for table in word_file.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
        if text.strip():
            documents.append(Document(
                page_content=text,
                metadata={"source": file_name, "file_type": "docx"}
            ))
    elif extension == ".xlsx":
        sheets = pd.read_excel(path, sheet_name=None)  
        for sheet_name, data in sheets.items():
            documents += table_to_documents(data, file_name, "xlsx", sheet_name)
    elif extension == ".csv":
        try:
            data = pd.read_csv(path, encoding="utf-8-sig")
        except UnicodeDecodeError:
            data = pd.read_csv(path, encoding="latin-1")
        documents += table_to_documents(data, file_name, "csv")
    elif extension == ".txt":
        try:
            with open(path, "r", encoding="utf-8-sig") as file:
                text = file.read()
        except UnicodeDecodeError:
            with open(path, "r", encoding="latin-1") as file:
                text = file.read()
        if text.strip():
            documents.append(Document(
                page_content=text,
                metadata={"source": file_name, "file_type": "txt"}
            ))
    else:
        raise ValueError(f"Unsupported file type: {extension}")
    if not documents:
        raise ValueError(
            f"No readable text found in {file_name}. "
            "If it is a scanned PDF, it needs OCR first."
        )
    return documents
def create_vectorstore(documents, collection_name=None):
    db = get_vectorstore(collection_name)
    file_names = {d.metadata.get("source") for d in documents if d.metadata.get("source")}
    for file_name in file_names:
        db.delete(where={"source": file_name})
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    chunks = splitter.split_documents(documents)
    if not chunks:
        raise ValueError("The file has no text to index.")
    try:
        for i in range(0, len(chunks), 50):
            batch = chunks[i:i + 50]
            for attempt in range(4):
                try:
                    db.add_documents(batch)
                    break
                except Exception:
                    if attempt == 3:
                        raise
                    time.sleep(2 * (attempt + 1))
    except Exception:
        for file_name in file_names:
            try:
                db.delete(where={"source": file_name})
            except Exception:
                pass
        raise
    return db
def delete_document(file_name, collection_name=None):
    db = get_vectorstore(collection_name)
    db.delete(where={"source": file_name})
    return True